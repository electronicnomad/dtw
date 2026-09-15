#!/usr/bin/env python3
import os
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SERVER_DIR = os.path.join(BASE_DIR, "server")
INDEX_HTML = os.path.join(SERVER_DIR, "index.html")
EXPORT_DIR = os.path.join(SERVER_DIR, "export_slides")
TOTAL_SLIDES = 29

CHROME_BIN = os.environ.get(
    "CHROME_BIN",
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    if os.path.exists("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    else "google-chrome"
)

os.makedirs(EXPORT_DIR, exist_ok=True)

def capture_slide(slide_num):
    out_path = os.path.join(EXPORT_DIR, f"slide_{slide_num:02d}.png")
    url = f"file://{INDEX_HTML}?export=true&slide={slide_num}"
    cmd = [
        CHROME_BIN,
        "--headless",
        "--disable-gpu",
        "--hide-scrollbars",
        "--window-size=1920,1080",
        f"--screenshot={out_path}",
        url
    ]
    result = subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if result.returncode == 0 and os.path.exists(out_path):
        print(f"Slide {slide_num:02d}/{TOTAL_SLIDES:02d} captured successfully.")
        return out_path
    else:
        print(f"Error capturing slide {slide_num:02d}", file=sys.stderr)
        return None

def main():
    # Remove old slide images in EXPORT_DIR
    for f in os.listdir(EXPORT_DIR):
        if f.endswith('.png'):
            os.remove(os.path.join(EXPORT_DIR, f))

    print("Capturing 29 presentation slides via headless Chrome...")
    with ThreadPoolExecutor(max_workers=4) as executor:
        futures = [executor.submit(capture_slide, i) for i in range(1, TOTAL_SLIDES + 1)]
        slide_paths = [f.result() for f in futures]
    
    # Sort paths by slide number
    slide_paths = [os.path.join(EXPORT_DIR, f"slide_{i:02d}.png") for i in range(1, TOTAL_SLIDES + 1)]
    
    # Check that all slides exist
    for p in slide_paths:
        if not os.path.exists(p):
            print(f"Missing slide: {p}", file=sys.stderr)
            sys.exit(1)
            
    print("\nGenerating 16:9 PDF presentation...")
    images = []
    for p in slide_paths:
        img = Image.open(p).convert("RGB")
        images.append(img)
        
    pdf_paths = [
        os.path.join(SERVER_DIR, "dtw_presentation.pdf"),
        os.path.join(BASE_DIR, "dtw_presentation.pdf")
    ]
    for pdf_out in pdf_paths:
        images[0].save(pdf_out, save_all=True, append_images=images[1:], resolution=150.0)
        print(f"PDF saved: {pdf_out} ({os.path.getsize(pdf_out) / (1024*1024):.2f} MB)")

    print("\nGenerating 16:9 PowerPoint (.pptx) presentation...")
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    for p in slide_paths:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(p, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)
        
    pptx_paths = [
        os.path.join(SERVER_DIR, "dtw_presentation.pptx"),
        os.path.join(BASE_DIR, "dtw_presentation.pptx")
    ]
    for pptx_out in pptx_paths:
        prs.save(pptx_out)
        print(f"PPTX saved: {pptx_out} ({os.path.getsize(pptx_out) / (1024*1024):.2f} MB)")

    print(f"\nExport completed successfully! All {TOTAL_SLIDES} slides exported to PDF & PPTX.")

if __name__ == "__main__":
    main()
